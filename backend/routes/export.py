"""
Exportación a PDF y Excel.
"""
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from io import BytesIO
from datetime import datetime

from ..database import get_db
from ..models import Student, AvacAccess, TaskSubmission, Grade, Intervention
from ..models.enrollment import Enrollment
from ..auth.jwt import get_current_user
from ..models.user import User

router = APIRouter(prefix="/export", tags=["export"])


# ─── Meses en español ────────────────────────────────────────────────────────

MESES_ES = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril",
    5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
    9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre",
}


# ─── Helper: Hoja LÉEME ──────────────────────────────────────────────────────

def _add_leeme_sheet(wb, report_title: str, filters_desc: str, user_name: str, generated_at: datetime):
    """
    Inserta una hoja 'LÉEME' como primera hoja del workbook con información
    de autoría, condiciones de uso y cita sugerida.
    """
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("LÉEME", 0)

    # ── Estilos ──────────────────────────────────────────────────────────────
    NAVY = "1B3A6B"
    GRAY_BG = "F5F7FA"
    GOLD = "F0B000"

    title_font = Font(name="Calibri", bold=True, color=NAVY, size=14)
    section_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    section_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    content_font = Font(name="Calibri", size=10)
    content_font_bold = Font(name="Calibri", size=10, bold=True)
    meta_font = Font(name="Calibri", size=10, color="555555")
    gold_font = Font(name="Calibri", size=10, bold=True, color=GOLD.replace("#", ""))
    thin_border = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    wrap_align = Alignment(vertical="top", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ── Ancho de columna ─────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 80

    # ── Fecha dinámica ───────────────────────────────────────────────────────
    dia = generated_at.day
    mes = MESES_ES.get(generated_at.month, str(generated_at.month))
    anio = generated_at.year
    fecha_str = f"{dia} de {mes} de {anio}"

    # ── Contenido ────────────────────────────────────────────────────────────
    row = 1

    def write_title(text):
        nonlocal row
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = title_font
        cell.alignment = center_align
        cell.border = thin_border
        row += 1

    def write_section(text):
        nonlocal row
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = section_font
        cell.fill = section_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = thin_border
        row += 1

    def write_line(text, font=None, indent=False):
        nonlocal row
        display = ("    " + text) if indent else text
        cell = ws.cell(row=row, column=1, value=display)
        cell.font = font or content_font
        cell.alignment = wrap_align
        cell.border = thin_border
        row += 1

    def write_blank():
        nonlocal row
        cell = ws.cell(row=row, column=1, value="")
        cell.border = thin_border
        row += 1

    # ════════════════════════════════════════════════════════════════════════
    # TÍTULO
    # ════════════════════════════════════════════════════════════════════════
    write_blank()
    write_title("DOCUMENTACIÓN Y TÉRMINOS DE USO DE DATOS")
    write_blank()

    # ── 1. INFORMACIÓN DE AUTORÍA ─────────────────────────────────────────
    write_section("1. INFORMACIÓN DE AUTORÍA Y PROPIEDAD INTELECTUAL")
    write_blank()
    write_line("• Desarrollado por:     Carlos Vásconez-Paredes", content_font_bold)
    write_line("• Cargo/Función:        Gestor de Analítica del Aprendizaje")
    write_line("• Institución:          Universidad Politécnica Salesiana")
    write_line(f"• Fecha de generación:  {fecha_str}")
    write_line("• Versión del dataset:  v1.0 (Estructurado y Procesado)")
    write_line(f"• Reporte:              {report_title}")
    write_blank()

    # ── 2. CONDICIONES DE USO ─────────────────────────────────────────────
    write_section("2. CONDICIONES DE USO Y RECONOCIMIENTO (LICENCIA)")
    write_blank()
    write_line(
        "Este conjunto de datos, métricas e interpretaciones analíticas son el resultado "
        "de un desarrollo metodológico y técnico específico. Se autoriza su uso para "
        "fines académicos, artículos científicos, ponencias y conferencias, bajo la "
        "condición estricta de otorgar el crédito correspondiente al autor."
    )
    write_blank()
    write_line(
        "De acuerdo con las políticas de integridad científica, la omisión de la fuente "
        "se considerará una falta a la ética académica."
    )
    write_blank()

    # ── 3. FORMA SUGERIDA DE CITA ─────────────────────────────────────────
    write_section("3. FORMA SUGERIDA DE CITA / REFERENCIA")
    write_blank()
    write_line("• Estilo APA (7ma ed.):", content_font_bold)
    write_line(
        f"  Vásconez-Paredes, C. ({anio}). {report_title} "
        f"(Versión 1.0) [Conjunto de datos/Métricas analíticas]. Gestión de Analítica "
        f"del Aprendizaje, Universidad Politécnica Salesiana.",
        indent=True,
    )
    write_blank()
    write_line("• Estilo Vancouver / Nota al pie:", content_font_bold)
    write_line(
        f"  Datos analíticos y procesamiento metodológico provistos por Carlos "
        f"Vásconez-Paredes, Gestión de Analítica del Aprendizaje, Universidad "
        f"Politécnica Salesiana, {anio}.",
        indent=True,
    )
    write_blank()

    # ── 4. CONTACTO Y COLABORACIÓN ────────────────────────────────────────
    write_section("4. CONTACTO Y COLABORACIÓN")
    write_blank()
    write_line(
        "Si su investigación requiere modificaciones metodológicas en los datos, cruces "
        "de variables avanzados o una interpretación analítica conjunta que impacte la "
        'sección de "Metodología" o "Resultados" del artículo, por favor tome contacto '
        "para estructurar una participación formal bajo la figura de coautoría."
    )
    write_blank()
    write_line("Contacto: cvasconez@ups.edu.ec", content_font_bold)
    write_blank()

    # ── 5. FILTROS APLICADOS ──────────────────────────────────────────────
    write_section("5. FILTROS APLICADOS EN ESTA EXPORTACIÓN")
    write_blank()
    write_line(filters_desc)
    write_blank()

    # ── Metadatos finales ─────────────────────────────────────────────────
    write_line(f"Generado por: {user_name}", meta_font)
    write_line("Plataforma:   YachayDeep — Sistema de Analítica del Aprendizaje", meta_font)
    write_blank()

    # ── Aplicar borde a toda el área usada ────────────────────────────────
    for r in range(1, row):
        cell = ws.cell(row=r, column=1)
        cell.border = thin_border

    return ws


# ─── Columnas disponibles para exportación Excel ─────────────────────────────

EXPORT_COLUMNS = {
    "cedula": {"label": "Cédula", "getter": lambda s, _: s.cedula},
    "nombre": {"label": "Nombres completos", "getter": lambda s, _: s.nombre},
    "correo": {"label": "Correo personal", "getter": lambda s, _: s.correo},
    "correo_institucional": {"label": "Correo institucional", "getter": lambda s, _: s.correo_institucional},
    "telefono": {"label": "Teléfono", "getter": lambda s, _: s.telefono},
    "whatsapp": {"label": "WhatsApp", "getter": lambda s, _: s.whatsapp},
    "carrera": {"label": "Carrera", "getter": lambda s, _: s.carrera},
    "nivel_academico": {"label": "Nivel académico", "getter": lambda s, _: s.nivel_academico},
    "sede": {"label": "Centro de apoyo", "getter": lambda s, _: s.sede},
    "grupo": {"label": "Grupo", "getter": lambda s, _: s.grupo},
    "estado_matricula": {"label": "Estado matrícula", "getter": lambda s, _: s.estado_matricula},
    "nivel_riesgo": {"label": "Nivel de riesgo", "getter": lambda s, _: s.nivel_riesgo},
    "indice_compromiso": {"label": "Índice compromiso", "getter": lambda s, _: round(s.indice_compromiso, 2) if s.indice_compromiso is not None else None},
    "dias_sin_acceso": {"label": "Días sin acceso AVAC", "getter": lambda s, _: s.dias_sin_acceso},
    "porcentaje_tareas": {"label": "% Tareas entregadas", "getter": lambda s, _: round(s.porcentaje_tareas, 1) if s.porcentaje_tareas is not None else None},
    "promedio_calificaciones": {"label": "Promedio calificaciones", "getter": lambda s, _: round(s.promedio_calificaciones, 1) if s.promedio_calificaciones is not None else None},
    "prob_desercion": {"label": "Prob. deserción", "getter": lambda s, _: round(s.prob_desercion, 2) if s.prob_desercion is not None else None},
    "prob_reprobacion": {"label": "Prob. reprobación", "getter": lambda s, _: round(s.prob_reprobacion, 2) if s.prob_reprobacion is not None else None},
    "genero": {"label": "Género", "getter": lambda s, _: s.genero},
    "autoidentificacion_etnica": {"label": "Autoidentificación étnica", "getter": lambda s, _: s.autoidentificacion_etnica},
    "fecha_nacimiento": {"label": "Fecha nacimiento", "getter": lambda s, _: s.fecha_nacimiento.isoformat() if s.fecha_nacimiento else None},
    "pais": {"label": "País", "getter": lambda s, _: s.pais},
    "provincia": {"label": "Provincia", "getter": lambda s, _: s.provincia},
    "ciudad": {"label": "Ciudad", "getter": lambda s, _: s.ciudad},
    "parroquia": {"label": "Parroquia", "getter": lambda s, _: s.parroquia},
    "barrio": {"label": "Barrio/Comunidad", "getter": lambda s, _: s.barrio},
    "total_intervenciones": {"label": "Total intervenciones", "getter": lambda s, ctx: ctx.get("interv", {}).get(s.id, 0)},
}


@router.get("/columnas-disponibles")
def get_columnas_disponibles(
    current_user: User = Depends(get_current_user),
):
    """Lista de columnas disponibles para exportación con su clave y etiqueta."""
    return [{"key": k, "label": v["label"]} for k, v in EXPORT_COLUMNS.items()]


@router.get("/estudiantes/excel")
def export_estudiantes_excel(
    carrera: Optional[str] = None,
    nivel: Optional[int] = None,
    nivel_riesgo: Optional[str] = None,
    periodo: Optional[str] = None,
    columnas: str = Query("cedula,nombre,correo_institucional,carrera,nivel_academico,nivel_riesgo", description="Columnas separadas por coma"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Exporta estudiantes a Excel con filtros y columnas seleccionables.
    periodo: 'actual' o None = semestre actual, 'P60'-'P67' = histórico, 'todos' = todos.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no instalado")

    from ..models import Grade as _Grade
    from sqlalchemy import or_

    periodo_filter = periodo if periodo else "actual"

    # Si es período histórico, filtrar estudiantes por quienes tienen calificaciones O matrículas en ese período
    student_ids_in_periodo = None
    if periodo_filter not in ("actual", "todos"):
        pf = periodo_filter
        # Dual format normalization: "P68" <-> "68"
        if pf.startswith("P"):
            grade_sids = db.query(_Grade.student_id).filter(or_(_Grade.periodo == pf, _Grade.periodo == pf[1:])).distinct().all()
            enroll_sids = db.query(Enrollment.student_id).filter(or_(Enrollment.periodo == pf, Enrollment.periodo == pf[1:])).distinct().all()
        else:
            grade_sids = db.query(_Grade.student_id).filter(or_(_Grade.periodo == pf, _Grade.periodo == f"P{pf}")).distinct().all()
            enroll_sids = db.query(Enrollment.student_id).filter(or_(Enrollment.periodo == pf, Enrollment.periodo == f"P{pf}")).distinct().all()

        student_ids_in_periodo = set(sid for (sid,) in grade_sids) | set(sid for (sid,) in enroll_sids)
        if not student_ids_in_periodo:
            raise HTTPException(status_code=404, detail=f"No hay estudiantes para el período {periodo_filter}")

    # Filtrar estudiantes
    query = db.query(Student)
    if carrera:
        query = query.filter(func.lower(Student.carrera).contains(carrera.lower()))
    if nivel:
        query = query.filter(Student.nivel_academico == nivel)
    if nivel_riesgo:
        query = query.filter(Student.nivel_riesgo == nivel_riesgo)
    if student_ids_in_periodo is not None:
        query = query.filter(Student.id.in_(student_ids_in_periodo))
    query = query.order_by(Student.carrera, Student.nivel_academico, Student.nombre)
    students = query.all()

    if not students:
        raise HTTPException(status_code=404, detail="No se encontraron estudiantes con los filtros aplicados")

    # Parsear columnas solicitadas
    cols_requested = [c.strip() for c in columnas.split(",") if c.strip() in EXPORT_COLUMNS]
    if not cols_requested:
        cols_requested = ["cedula", "nombre", "correo_institucional", "carrera", "nivel_academico", "nivel_riesgo"]

    # Pre-cargar contexto (intervenciones por estudiante)
    ctx = {}
    if "total_intervenciones" in cols_requested:
        student_ids = [s.id for s in students]
        interv_counts = dict(
            db.query(Intervention.student_id, func.count(Intervention.id))
            .filter(Intervention.student_id.in_(student_ids))
            .group_by(Intervention.student_id)
            .all()
        )
        ctx["interv"] = interv_counts

    # Crear workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Estudiantes"

    # Estilos
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1B3A6B", end_color="1B3A6B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    alt_fill = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")

    # Headers
    for col_idx, col_key in enumerate(cols_requested, 1):
        cell = ws.cell(row=1, column=col_idx, value=EXPORT_COLUMNS[col_key]["label"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Datos
    for row_idx, student in enumerate(students, 2):
        for col_idx, col_key in enumerate(cols_requested, 1):
            value = EXPORT_COLUMNS[col_key]["getter"](student, ctx)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    # Autofit columns
    for col_idx, col_key in enumerate(cols_requested, 1):
        max_len = len(EXPORT_COLUMNS[col_key]["label"])
        for row_idx in range(2, min(len(students) + 2, 52)):  # Sample first 50 rows
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 3, 40)

    # Fila de resumen
    summary_row = len(students) + 3
    ws.cell(row=summary_row, column=1, value=f"Total: {len(students)} estudiantes").font = Font(bold=True, size=10)
    ws.cell(row=summary_row + 1, column=1, value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')} — {current_user.nombre}").font = Font(size=9, color="888888")

    # Filtro aplicado
    filter_desc = "Filtros: "
    if periodo_filter and periodo_filter != "actual":
        filter_desc += f"Período={periodo_filter} "
    else:
        filter_desc += "Período=Semestre actual "
    if carrera:
        filter_desc += f"Carrera={carrera} "
    if nivel:
        filter_desc += f"Nivel={nivel} "
    if nivel_riesgo:
        filter_desc += f"Riesgo={nivel_riesgo} "
    if not carrera and not nivel and not nivel_riesgo:
        filter_desc += "(sin filtros adicionales)"
    ws.cell(row=summary_row + 2, column=1, value=filter_desc).font = Font(size=9, color="888888")

    # Freeze header row
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ── Agregar hoja LÉEME ────────────────────────────────────────────────
    now = datetime.now()
    report_title_str = f"Estudiantes — {carrera}" if carrera else "Estudiantes — Todas las carreras"
    _add_leeme_sheet(wb, report_title_str, filter_desc, current_user.nombre, now)

    # Asegurar que la hoja de datos sea la activa al abrir
    wb.active = wb.sheetnames.index("Estudiantes")

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    import re
    safe_carrera = re.sub(r'[^\w\s-]', '', carrera or "todos").replace(' ', '_')
    filename = f"estudiantes_{safe_carrera}_{datetime.now().strftime('%Y%m%d')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ─── Columnas disponibles para exportación de Intervenciones ──────────────────

INTERVENTION_COLUMNS = {
    "nombre": {"label": "Estudiante", "getter": lambda inv, ctx: ctx["students"].get(inv.student_id, {}).get("nombre")},
    "cedula": {"label": "Cédula", "getter": lambda inv, ctx: ctx["students"].get(inv.student_id, {}).get("cedula")},
    "correo_institucional": {"label": "Correo institucional", "getter": lambda inv, ctx: ctx["students"].get(inv.student_id, {}).get("correo_institucional")},
    "telefono": {"label": "Teléfono", "getter": lambda inv, ctx: ctx["students"].get(inv.student_id, {}).get("telefono")},
    "whatsapp": {"label": "WhatsApp", "getter": lambda inv, ctx: ctx["students"].get(inv.student_id, {}).get("whatsapp")},
    "carrera": {"label": "Carrera", "getter": lambda inv, ctx: ctx["students"].get(inv.student_id, {}).get("carrera") or inv.carrera},
    "sede": {"label": "Centro de apoyo", "getter": lambda inv, ctx: ctx["students"].get(inv.student_id, {}).get("sede")},
    "nivel_riesgo": {"label": "Nivel de riesgo", "getter": lambda inv, ctx: ctx["students"].get(inv.student_id, {}).get("nivel_riesgo")},
    "medio": {"label": "Medio de contacto", "getter": lambda inv, _: inv.medio},
    "motivo": {"label": "Motivo", "getter": lambda inv, _: inv.motivo},
    "estado": {"label": "Estado", "getter": lambda inv, _: inv.estado},
    "resultado": {"label": "Resultado", "getter": lambda inv, _: inv.resultado},
    "asignatura": {"label": "Asignatura", "getter": lambda inv, _: inv.asignatura},
    "docente": {"label": "Docente", "getter": lambda inv, _: inv.docente},
    "observacion": {"label": "Observación", "getter": lambda inv, _: inv.observacion},
    "requiere_seguimiento": {"label": "Requiere seguimiento", "getter": lambda inv, _: "Sí" if inv.requiere_seguimiento == "si" else ("No" if inv.requiere_seguimiento == "no" else "")},
    "derivar_bienestar": {"label": "Derivado a Bienestar", "getter": lambda inv, _: "Sí" if inv.derivar_bienestar else "No"},
    "tipo_evento_critico": {"label": "Tipo evento crítico", "getter": lambda inv, _: inv.tipo_evento_critico},
    "reporte_bienestar": {"label": "Reporte Bienestar", "getter": lambda inv, _: inv.reporte_bienestar},
    "email_enviado": {"label": "Email Bienestar enviado", "getter": lambda inv, _: "Sí" if inv.email_enviado else "No"},
    "monitor_nombre": {"label": "Monitor", "getter": lambda inv, _: inv.monitor_nombre},
    "periodo": {"label": "Período", "getter": lambda inv, _: inv.periodo},
    "fecha": {"label": "Fecha", "getter": lambda inv, _: inv.created_at.strftime("%d/%m/%Y %H:%M") if inv.created_at else None},
}


@router.get("/intervenciones/columnas-disponibles")
def get_columnas_intervenciones(
    current_user: User = Depends(get_current_user),
):
    """Lista de columnas disponibles para exportación de intervenciones."""
    return [{"key": k, "label": v["label"]} for k, v in INTERVENTION_COLUMNS.items()]


@router.get("/intervenciones/excel")
def export_intervenciones_excel(
    carrera: Optional[str] = None,
    motivo: Optional[str] = None,
    estado: Optional[str] = None,
    resultado: Optional[str] = None,
    seguimiento: Optional[str] = None,
    columnas: str = Query(
        "nombre,carrera,nivel_riesgo,medio,motivo,estado,resultado,requiere_seguimiento,observacion,monitor_nombre,fecha",
        description="Columnas separadas por coma",
    ),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Exporta intervenciones a Excel con filtros y columnas seleccionables."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no instalado")

    # Query con los mismos filtros que el dashboard
    query = db.query(Intervention).join(Student, Intervention.student_id == Student.id)
    if carrera:
        query = query.filter(Student.carrera == carrera)
    if motivo:
        query = query.filter(Intervention.motivo == motivo)
    if estado:
        query = query.filter(Intervention.estado == estado)
    if resultado:
        query = query.filter(Intervention.resultado == resultado)
    if seguimiento:
        query = query.filter(Intervention.requiere_seguimiento == seguimiento)

    interventions = query.order_by(Intervention.created_at.desc()).all()

    if not interventions:
        raise HTTPException(status_code=404, detail="No se encontraron intervenciones con los filtros aplicados")

    # Parsear columnas
    cols_requested = [c.strip() for c in columnas.split(",") if c.strip() in INTERVENTION_COLUMNS]
    if not cols_requested:
        cols_requested = ["nombre", "carrera", "motivo", "estado", "resultado", "monitor_nombre", "fecha"]

    # Pre-cargar datos de estudiantes
    student_ids = list(set(inv.student_id for inv in interventions))
    students_data = db.query(Student).filter(Student.id.in_(student_ids)).all()
    ctx = {
        "students": {
            s.id: {
                "nombre": s.nombre,
                "cedula": s.cedula,
                "correo_institucional": s.correo_institucional,
                "telefono": s.telefono,
                "whatsapp": s.whatsapp,
                "carrera": s.carrera,
                "sede": s.sede,
                "nivel_riesgo": s.nivel_riesgo,
            }
            for s in students_data
        }
    }

    # Crear workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Intervenciones"

    # Estilos
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1B3A6B", end_color="1B3A6B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    alt_fill = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")

    # Headers
    for col_idx, col_key in enumerate(cols_requested, 1):
        cell = ws.cell(row=1, column=col_idx, value=INTERVENTION_COLUMNS[col_key]["label"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Datos
    for row_idx, inv in enumerate(interventions, 2):
        for col_idx, col_key in enumerate(cols_requested, 1):
            value = INTERVENTION_COLUMNS[col_key]["getter"](inv, ctx)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    # Autofit columns
    for col_idx, col_key in enumerate(cols_requested, 1):
        max_len = len(INTERVENTION_COLUMNS[col_key]["label"])
        for row_idx in range(2, min(len(interventions) + 2, 52)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 3, 50)

    # Resumen
    summary_row = len(interventions) + 3
    ws.cell(row=summary_row, column=1, value=f"Total: {len(interventions)} intervenciones").font = Font(bold=True, size=10)
    ws.cell(row=summary_row + 1, column=1, value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')} — {current_user.nombre}").font = Font(size=9, color="888888")

    filter_parts = ["Filtros:"]
    if carrera:
        filter_parts.append(f"Carrera={carrera}")
    if motivo:
        filter_parts.append(f"Motivo={motivo}")
    if estado:
        filter_parts.append(f"Estado={estado}")
    if resultado:
        filter_parts.append(f"Resultado={resultado}")
    if seguimiento:
        filter_parts.append(f"Seguimiento={seguimiento}")
    if len(filter_parts) == 1:
        filter_parts.append("(sin filtros)")
    filter_desc = " ".join(filter_parts)
    ws.cell(row=summary_row + 2, column=1, value=filter_desc).font = Font(size=9, color="888888")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ── Agregar hoja LÉEME ────────────────────────────────────────────────
    now = datetime.now()
    report_title_str = f"Intervenciones — {carrera}" if carrera else "Intervenciones — Todas las carreras"
    _add_leeme_sheet(wb, report_title_str, filter_desc, current_user.nombre, now)

    # Asegurar que la hoja de datos sea la activa al abrir
    wb.active = wb.sheetnames.index("Intervenciones")

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    import re
    safe_carrera = re.sub(r'[^\w\s-]', '', carrera or "todas").replace(' ', '_')
    filename = f"intervenciones_{safe_carrera}_{datetime.now().strftime('%Y%m%d')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/ficha/{student_id}/pdf")
def export_ficha_pdf(
    student_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Genera PDF completo de la ficha del estudiante.
    Incluye: datos personales, residencia, estado socioeconómico,
    indicadores, calificaciones, actividades AVAC, intervenciones y prácticas.
    """
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Table, TableStyle,
            Spacer, HRFlowable, KeepTogether
        )
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    except ImportError:
        raise HTTPException(status_code=500, detail="reportlab no instalado")

    from ..models.practica_preprofesional import PracticaPreprofesional, EscuelaPractica

    student = db.query(Student).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Estudiante no encontrado")

    tareas = db.query(TaskSubmission).filter(TaskSubmission.student_id == student_id).all()
    accesos = db.query(AvacAccess).filter(AvacAccess.student_id == student_id).all()
    calificaciones = db.query(Grade).filter(Grade.student_id == student_id).order_by(Grade.periodo.desc()).all()
    intervenciones = db.query(Intervention).filter(Intervention.student_id == student_id).order_by(Intervention.created_at.desc()).all()
    practicas = (
        db.query(PracticaPreprofesional)
        .filter(PracticaPreprofesional.student_id == student_id)
        .all()
    )

    # ── Colores y estilos ──
    BRAND = colors.HexColor("#1B3A6B")
    BRAND_LIGHT = colors.HexColor("#D6E4F0")
    GOLD = colors.HexColor("#F0B000")
    GRAY_BG = colors.HexColor("#F5F7FA")
    GRAY_BORDER = colors.HexColor("#D0D5DD")

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        rightMargin=1.2 * cm, leftMargin=1.2 * cm,
        topMargin=1.2 * cm, bottomMargin=1.2 * cm,
    )
    styles = getSampleStyleSheet()
    story = []

    # ── Helpers ──
    section_style = ParagraphStyle(
        "section", fontSize=10, fontName="Helvetica-Bold",
        textColor=BRAND, spaceBefore=6, spaceAfter=3,
    )
    cell_p = lambda txt, bold=False, size=8, color=colors.black: Paragraph(
        f"<font size={size}><b>{txt}</b></font>" if bold
        else f"<font size={size} color='{color}'>{txt}</font>",
        styles["Normal"],
    )
    label_style = TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (0, -1), BRAND_LIGHT),
        ("GRID", (0, 0), (-1, -1), 0.5, GRAY_BORDER),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ])
    header_table_style = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BRAND),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, GRAY_BORDER),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, GRAY_BG]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ])

    v = lambda x: str(x) if x not in (None, "", "nan") else "—"
    pct = lambda x: f"{round(x * 100)}%" if x is not None else "—"

    # ════════════════════ HEADER ════════════════════
    title_style = ParagraphStyle(
        "title", fontSize=14, fontName="Helvetica-Bold",
        alignment=TA_CENTER, textColor=BRAND,
    )
    story.append(Paragraph("YACHAY DEEP — Ficha del Estudiante", title_style))
    story.append(Spacer(1, 0.2 * cm))
    story.append(HRFlowable(width="100%", thickness=2, color=BRAND))
    story.append(Spacer(1, 0.3 * cm))

    # ════════════════════ DATOS PERSONALES + RESIDENCIA (side by side) ════════════════════
    # Left: personal info
    edad = ""
    if student.fecha_nacimiento:
        from datetime import date
        born = student.fecha_nacimiento
        today = date.today()
        age = today.year - born.year - ((today.month, today.day) < (born.month, born.day))
        fecha_str = born.strftime("%d/%m/%Y")
        edad = f"{fecha_str} ({age} años)"

    personal_data = [
        ["Nombre", v(student.nombre)],
        ["Cédula", v(student.cedula)],
        ["Correo", v(student.correo)],
        ["Correo Institucional", v(student.correo_institucional)],
        ["Teléfono / WhatsApp", v(student.whatsapp or student.telefono)],
        ["Fecha nac. y edad", edad or "—"],
        ["Autoidentificación", v(student.autoidentificacion_etnica)],
        ["Género", v(student.genero)],
    ]
    t_personal = Table(personal_data, colWidths=[3.5 * cm, 7 * cm])
    t_personal.setStyle(label_style)

    # Right: residence + socioeconomic
    residence_data = [
        ["Provincia", v(student.provincia)],
        ["Cantón", v(student.ciudad)],
        ["Parroquia", v(student.parroquia)],
        ["Barrio", v(student.barrio)],
    ]
    matricula_text = "Ya pagó la matrícula" if student.estado_matricula == "Matriculado" else (
        "Aún no paga" if student.estado_matricula else "—"
    )
    socio_data = [
        ["Pago matrícula", matricula_text],
    ]
    right_data = residence_data + socio_data
    t_right = Table(right_data, colWidths=[3 * cm, 7.5 * cm])
    t_right.setStyle(label_style)

    # Wrap in a 2-column layout
    story.append(Paragraph("Datos personales / Residencia / Estado socioeconómico", section_style))
    layout = Table([[t_personal, t_right]], colWidths=[10.5 * cm, 10.5 * cm])
    layout.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (0, -1), 8),
    ]))
    story.append(layout)
    story.append(Spacer(1, 0.3 * cm))

    # ════════════════════ DATOS ACADÉMICOS ════════════════════
    story.append(Paragraph("Datos académicos", section_style))
    acad_data = [
        ["Carrera", v(student.carrera)],
        ["Centro de apoyo", v(student.sede)],
        ["Nivel", v(student.nivel_academico)],
        ["Grupo", v(student.grupo)],
    ]
    t_acad = Table(acad_data, colWidths=[3.5 * cm, 17.5 * cm])
    t_acad.setStyle(label_style)
    story.append(t_acad)
    story.append(Spacer(1, 0.3 * cm))

    # ════════════════════ INDICADORES ════════════════════
    story.append(Paragraph("Indicadores · Predicción IA", section_style))
    ind_data = [
        ["Indicador", "Valor", "Nivel"],
        [
            "Compromiso",
            pct(student.indice_compromiso),
            "Alto" if (student.indice_compromiso or 0) >= 0.6 else "Medio" if (student.indice_compromiso or 0) >= 0.3 else "Bajo",
        ],
        [
            "Predicción Deserción",
            pct(student.prob_desercion),
            "Alto" if (student.prob_desercion or 0) >= 0.7 else "Moderado" if (student.prob_desercion or 0) >= 0.4 else "Bajo",
        ],
        [
            "Predicción Reprobación",
            pct(student.prob_reprobacion),
            "Alto" if (student.prob_reprobacion or 0) >= 0.7 else "Moderado" if (student.prob_reprobacion or 0) >= 0.4 else "Bajo",
        ],
        [
            "Días sin AVAC",
            f"{round(student.dias_sin_acceso)}d" if student.dias_sin_acceso is not None else "—",
            "",
        ],
        [
            "Tareas entregadas",
            f"{round(student.porcentaje_tareas)}%" if student.porcentaje_tareas is not None else "—",
            "",
        ],
    ]
    t_ind = Table(ind_data, colWidths=[5 * cm, 4 * cm, 4 * cm])
    t_ind.setStyle(header_table_style)
    story.append(t_ind)
    story.append(Spacer(1, 0.3 * cm))

    # ════════════════════ CALIFICACIONES ════════════════════
    if calificaciones:
        story.append(Paragraph("Calificaciones", section_style))
        cal_headers = ["Período", "Asignatura", "Docente", "Nota Final", "Estado", "Repitencias"]
        cal_rows = [cal_headers]
        for g in calificaciones[:50]:
            estado = "Aprobado" if (g.nota_final or 0) >= 70 else "Reprobado" if g.nota_final is not None else "—"
            cal_rows.append([
                v(g.periodo), v(g.asignatura), v(g.docente),
                str(round(g.nota_final, 1)) if g.nota_final is not None else "—",
                estado,
                str(g.numero_repitencias) if g.numero_repitencias else "—",
            ])
        t_cal = Table(cal_rows, colWidths=[2.5 * cm, 8 * cm, 4 * cm, 2.5 * cm, 2.5 * cm, 2 * cm])
        t_cal.setStyle(header_table_style)
        # Color cells for reprobado
        for i, g in enumerate(calificaciones[:50], 1):
            if g.nota_final is not None and g.nota_final < 70:
                t_cal.setStyle(TableStyle([
                    ("TEXTCOLOR", (3, i), (4, i), colors.HexColor("#DC2626")),
                ]))
        story.append(t_cal)
        story.append(Spacer(1, 0.3 * cm))

    # ════════════════════ ACTIVIDADES AVAC ════════════════════
    if tareas:
        story.append(Paragraph("Actividades AVAC", section_style))
        act_headers = ["Curso", "Unidad", "Estado", "Calificación", "Entregada"]
        act_rows = [act_headers]
        for t in tareas[:40]:
            act_rows.append([
                v(t.codigo_curso), v(t.unidad),
                (t.estado or "—")[:45],
                str(t.calificacion) if t.calificacion is not None else "—",
                "Sí" if t.entregada else "No",
            ])
        t_act = Table(act_rows, colWidths=[4 * cm, 2.5 * cm, 9 * cm, 3 * cm, 2.5 * cm])
        t_act.setStyle(header_table_style)
        story.append(t_act)
        story.append(Spacer(1, 0.3 * cm))

    # ════════════════════ PRÁCTICAS PREPROFESIONALES ════════════════════
    if practicas:
        story.append(Paragraph("Prácticas Preprofesionales", section_style))
        for p in practicas:
            escuela = db.query(EscuelaPractica).filter(EscuelaPractica.id == p.escuela_id).first() if p.escuela_id else None
            prac_data = [
                ["Práctica", v(p.nombre_practica)],
                ["IE Práctica", v(p.nombre_escuela)],
                ["Ubicación", v(p.ubicacion_escuela)],
                ["AMIE", v(p.amie_escuela)],
                ["Período", v(p.periodo)],
            ]
            if escuela:
                jurisdiccion = escuela.jurisdiccion or "—"
                prac_data.append(["Jurisdicción", jurisdiccion])
            t_prac = Table(prac_data, colWidths=[3.5 * cm, 17.5 * cm])
            t_prac.setStyle(label_style)
            story.append(t_prac)
            story.append(Spacer(1, 0.15 * cm))
        story.append(Spacer(1, 0.2 * cm))

    # ════════════════════ INTERVENCIONES ════════════════════
    if intervenciones:
        story.append(Paragraph(f"Historial de Intervenciones ({len(intervenciones)})", section_style))
        int_headers = ["Fecha", "Monitor", "Medio", "Motivo", "Estado", "Resultado", "Observación"]
        int_rows = [int_headers]
        for i in intervenciones[:30]:
            int_rows.append([
                i.created_at.strftime("%d/%m/%Y") if i.created_at else "—",
                (i.monitor_nombre or "—")[:18],
                v(i.medio),
                v(i.motivo),
                v(i.estado),
                v(i.resultado),
                (i.observacion or "—")[:50],
            ])
        t_int = Table(int_rows, colWidths=[2.2 * cm, 3 * cm, 2.2 * cm, 3 * cm, 2.2 * cm, 3 * cm, 5.5 * cm])
        t_int.setStyle(header_table_style)
        story.append(t_int)

    # ════════════════════ FOOTER ════════════════════
    story.append(Spacer(1, 0.5 * cm))
    story.append(HRFlowable(width="100%", thickness=1, color=GRAY_BORDER))
    footer_style = ParagraphStyle("footer", fontSize=7, textColor=colors.grey, alignment=TA_CENTER)
    story.append(Paragraph(
        f"Generado por YachayDeep — {datetime.now().strftime('%d/%m/%Y %H:%M')} — {current_user.nombre}",
        footer_style,
    ))

    doc.build(story)
    buffer.seek(0)

    import re
    safe_name = re.sub(r'[^\w\s-]', '', student.nombre or "estudiante").replace(' ', '_')
    filename = f"ficha_{safe_name}_{datetime.now().strftime('%Y%m%d')}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
